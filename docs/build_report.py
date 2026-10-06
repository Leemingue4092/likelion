#!/usr/bin/env python3
"""실험 4 적분기 문안. 제출 파일은 docs/hwp/build.sh가 만드는 HWP다."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips
from pathlib import Path

ROOT = Path(__file__).resolve().parent
FONT = "NanumGothic"
BLUE = RGBColor(0x1E, 0x5A, 0xA0)
GRAY = RGBColor(0x55, 0x55, 0x55)
BLACK = RGBColor(0x00, 0x00, 0x00)


def set_run_font(run, size, bold=False, color=None):
    run.font.name = FONT
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color or BLACK
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), FONT)


def add_text(p, text, size=11, bold=False, color=None):
    run = p.add_run(text)
    set_run_font(run, size, bold, color)
    return run


def paragraph(doc, text, size=11, bold=False, color=None, align="left",
              before=0, after=2, line=1.15):
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if text:
        add_text(p, text, size, bold, color)
    return p


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_border(cell, color="000000", sz="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def set_cell_margins(cell, mm=1.2):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    twip = str(int(mm * 56.7))
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), twip)
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def clear_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def set_table_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = sum(widths)
    tblW.set(qn("w:w"), str(int(total * 56.7)))
    tblW.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w * 56.7)))
        grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(widths[i] * 56.7)))
            tcW.set(qn("w:type"), "dxa")


def center_table(table):
    tblPr = table._tbl.tblPr
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblPr.append(jc)


def write_cell(cell, text, size=11, bold=False, align="left", fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    add_text(p, text, size, bold)
    set_cell_border(cell)
    set_cell_margins(cell)
    if fill:
        shade_cell(cell, fill)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    center_table(table)
    set_table_widths(table, widths)
    for i, h in enumerate(headers):
        write_cell(table.rows[0].cells[i], h, size=11, bold=True, align="center", fill="E8EEF4")
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            align = "center" if c == 0 else "left"
            write_cell(table.rows[r + 1].cells[c], val, size=11, align=align)
    paragraph(doc, "", size=6, after=4)
    return table


def add_sentences(doc, lines, before_first=0):
    for i, line in enumerate(lines):
        paragraph(doc, line, size=11, before=before_first if i == 0 else 0, after=2)


def heading(doc, text):
    paragraph(doc, text, size=13, bold=True, before=12, after=6)


def subhead(doc, text):
    paragraph(doc, text, size=12, bold=True, before=10, after=4)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # 표지
    paragraph(doc, "순천향대학교", size=20, bold=True, color=BLUE,
              align="center", before=72, after=2)
    paragraph(doc, "공과대학  ·  회로이론실습설계 2", size=12, color=GRAY,
              align="center", before=0, after=8)

    logo = paragraph(doc, "", align="center", before=2, after=2)
    run = logo.add_run()
    run.add_picture(str(ROOT / "sch_logo.png"), width=Mm(24))

    paragraph(doc, "실험 보고서", size=22, bold=True, color=BLUE,
              align="center", before=8, after=6)
    paragraph(doc, "실험 4.  적분기", size=16, bold=True,
              align="center", before=2, after=12)

    meta = [
        ("과목명", "회로이론실습설계 2"),
        ("교수명", "강병권 교수님"),
        ("학번", "20234092"),
        ("제출자", "이민규"),
        ("제출일", "입력 없음"),
        ("실험일", "입력 없음"),
        ("조 / 조원", ""),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    center_table(table)
    clear_table_borders(table)
    set_table_widths(table, [32, 78])
    for i, (label, value) in enumerate(meta):
        left, right = table.rows[i].cells
        left.text = ""
        right.text = ""
        pl = left.paragraphs[0]
        pl.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pl.paragraph_format.space_before = Pt(1)
        pl.paragraph_format.space_after = Pt(1)
        add_text(pl, label, 12)
        pr = right.paragraphs[0]
        pr.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pr.paragraph_format.space_before = Pt(1)
        pr.paragraph_format.space_after = Pt(1)
        add_text(pr, ": " + value if value else ":", 12)
        set_cell_margins(left, 0.4)
        set_cell_margins(right, 0.4)

    paragraph(doc, "순천향대학교  ·  2026학년도 2학기", size=11, color=GRAY,
              align="center", before=16, after=0)

    doc.add_page_break()

    heading(doc, "1. 실험 목적")
    add_sentences(doc, [
        "적분기 회로에 사인파와 삼각파를 넣어 출력 모양을 보는 실험이다.",
        "출력은 입력을 시간에 대해 더한 값에 비례한다.",
        "수업 화면에는 10 Vpp, 3 kHz, 사인파, 삼각파가 적혀 있다.",
    ])

    heading(doc, "2. 실험 이론")
    subhead(doc, "2-1. 적분기가 하는 일")
    add_sentences(doc, [
        "출력은 입력의 적분에 비례한다.",
        "사인 입력은 이렇다.",
    ])
    paragraph(doc, "Vi = Vp sin(ωt)", size=11, before=2, after=2)
    add_sentences(doc, [
        "이 회로는 입력을 핀 2, 반전 −단자에 넣는다.",
    ])
    eq = paragraph(doc, "", size=11, before=2, after=2)
    eq.paragraph_format.tab_stops.add_tab_stop(Mm(170), WD_TAB_ALIGNMENT.RIGHT)
    add_text(eq, "Vo = − (1 / (R C)) ∫ Vi dt")
    add_text(eq, "\t수식 2-1")
    add_sentences(doc, [
        "sin(ωt)의 적분은 −cos(ωt) / ω다.",
        "식 앞의 마이너스가 그 부호를 뒤집는다.",
    ])
    paragraph(doc, "Vo = (Vp / (ω R C)) cos(ωt)", size=11, before=2, after=2)
    add_sentences(doc, [
        "cos(ωt)는 sin(ωt)보다 90도 앞선 파다.",
        "출력은 코사인 모양이다.",
        "커패시터만 두면 출력은 입력보다 90도 앞선다.",
        "스코프에서 어느 채널이 어느 쪽으로 밀렸는지는 입력 없음.",
        "핀 3(+)는 접지다.",
        "핀 2는 거의 0 V다.",
        "입력이 양이면 15 kΩ으로 핀 2 쪽에 전류가 흐른다.",
        "그 전류는 핀 안으로 들어가지 못한다.",
        "그 전류는 100 nF를 타고 출력에서 나온다.",
        "그래서 출력은 아래로 간다.",
        "삼각파가 0보다 클 때 출력은 내려간다.",
        "삼각파가 0보다 작을 때 출력은 올라간다.",
        "직선인 입력을 적분하면 시간 제곱에 비례한다.",
        "삼각파 출력은 포물선이다.",
    ])

    subhead(doc, "2-2. 이번 회로 값")
    add_sentences(doc, [
        "입력 저항은 15 kΩ이다.",
        "피드백은 33 kΩ과 100 nF의 병렬이다.",
        "직류에서 100 nF는 열린다.",
        "이때 피드백은 33 kΩ만 남는다.",
        "직류 배율은 −33 kΩ / 15 kΩ = −2.2다.",
        "Vp는 10 Vpp의 절반이라 5 V다.",
        "ω R C = 2π · 3 kHz · 15 kΩ · 100 nF = 28.27이다.",
        "1 / 28.27 = 0.03537이다.",
        "사인파 이론 출력은 10 Vpp · 0.03537 = 0.3537 Vpp다.",
        "이 값은 15 kΩ과 100 nF만 쓴 계산이다.",
        "2π · 3 kHz · 33 kΩ · 100 nF = 62.20이다.",
        "1 + 62.20의 제곱의 제곱근은 62.21이다.",
        "2.2 / 62.21 = 0.03536이다.",
        "33 kΩ을 병렬로 넣은 사인파 출력은 10 Vpp · 0.03536 = 0.3536 Vpp다.",
        "반전 입력의 부호는 180도다.",
        "arctan(62.20) = 89.08도다.",
        "180 − 89.08 = 90.92도다.",
        "33 kΩ을 병렬로 넣은 사인파 출력은 입력보다 90.92도 앞선다.",
        "화면에서 읽은 입력으로 다시 계산한 값은 입력 없음.",
        "삼각파도 칠판의 전원 표시 10 Vpp로 계산했다.",
        "5 V / (4 · 3 kHz) = 4.167×10⁻⁴ V·s다.",
        "1 / (15 kΩ · 100 nF) = 666.7 s⁻¹이다.",
        "666.7 · 4.167×10⁻⁴ = 0.2778 Vpp다.",
        "삼각파 이론 출력의 pk-pk는 0.2778 Vpp다.",
        "이 값은 15 kΩ과 100 nF만 쓴 계산이다.",
    ])

    subhead(doc, "2-3. 회로도")
    paragraph(
        doc,
        "입력은 15 kΩ을 거쳐 핀 2(−)로 들어간다. 핀 3(+)는 접지다. "
        "33 kΩ과 100 nF는 핀 6과 핀 2 사이에 병렬로 있다. 핀 7은 +15 V이고 핀 4는 −15 V다.",
        size=11, before=0, after=6,
    )
    pic = paragraph(doc, "", align="center", before=2, after=2)
    pic.add_run().add_picture(str(ROOT / "수업화면_실험4_적분기.jpg"), width=Mm(70))
    paragraph(doc, "[그림 2-1] 수업 화면. 실험 4 적분기", size=10,
              align="center", before=2, after=8)

    subhead(doc, "2-4. 핀 정리")
    add_sentences(doc, [
        "칠판에 적힌 핀은 이렇다.",
    ])
    paragraph(doc, "[표 2-1] 핀 정리", size=10, align="center", before=4, after=2)
    add_table(
        doc,
        ["핀", "연결"],
        [
            ["핀 3 (+)", "접지"],
            ["핀 2 (−)", "입력 15 kΩ"],
            ["핀 6 (출력)", "Vo, 피드백 33 kΩ과 100 nF"],
            ["핀 7", "+15 V"],
            ["핀 4", "−15 V"],
        ],
        [40, 120],
    )
    add_sentences(doc, [
        "칠판에 없는 핀은 입력 없음.",
    ])

    heading(doc, "3. 실험기기 및 부품")
    paragraph(doc, "[표 3-1] 실험기기 및 부품", size=10, align="center", before=2, after=2)
    add_table(
        doc,
        ["항목", "내용"],
        [
            ["능동소자", "연산증폭기 (모델명 입력 없음)"],
            ["저항", "15 kΩ (입력), 33 kΩ (피드백)"],
            ["커패시터 또는 인덕터", "100 nF"],
            ["입력(함수발생기 설정)", "3 kHz, 10 Vpp, 사인파 / 삼각파"],
            ["측정기", "입력 없음"],
        ],
        [52, 108],
    )

    heading(doc, "4. 실험 방법")
    subhead(doc, "4-1. 주의사항")
    add_sentences(doc, [
        "핀 7은 +15 V다.",
        "핀 4는 −15 V다.",
        "입력은 15 kΩ을 거쳐 핀 2(−)로 넣는다.",
        "프로브 배율은 입력 없음.",
    ])

    subhead(doc, "4-2. 회로 구성")
    add_sentences(doc, [
        "브레드보드에 꽂은 순서는 입력 없음.",
        "스코프 CH1, CH2가 어느 노드인지는 입력 없음.",
    ])

    subhead(doc, "4-3. 측정 조건")
    add_sentences(doc, [
        "함수발생기 모델은 입력 없음.",
        "칠판 주파수는 3 kHz다.",
        "파형은 사인파와 삼각파다.",
        "진폭은 10 Vpp다.",
        "오프셋은 입력 없음.",
        "스코프 모델은 입력 없음.",
        "시간축은 입력 없음.",
        "전압축은 입력 없음.",
        "커플링은 입력 없음.",
        "프로브 배율은 입력 없음.",
        "촬영 시각은 입력 없음.",
    ])

    heading(doc, "5. 실험 결과")
    subhead(doc, "5-1. 사인파")
    add_sentences(doc, [
        "칠판에 적힌 첫 입력은 사인파다.",
        "진폭은 10 Vpp다.",
        "주파수는 3 kHz다.",
        "스코프 화면의 색, 채널, pk-pk, 시간축, 전압축, 촬영 시각은 입력 없음.",
        "15 kΩ과 100 nF만 쓴 이론 출력은 0.3537 Vpp다.",
        "33 kΩ을 병렬로 넣은 이론 출력은 0.3536 Vpp다.",
    ])
    paragraph(doc, "[표 5-1] 사인파", size=10, align="center", before=4, after=2)
    add_table(
        doc,
        ["항목", "값"],
        [
            ["칠판 입력", "사인파, 10 Vpp, 3 kHz"],
            ["스코프 pk-pk", "입력 없음"],
            ["시간축", "입력 없음"],
            ["전압축", "입력 없음"],
            ["이론 출력 (15 kΩ, 100 nF)", "0.3537 Vpp (계산)"],
            ["이론 출력 (33 kΩ 병렬)", "0.3536 Vpp (계산)"],
        ],
        [62, 98],
    )

    subhead(doc, "5-2. 삼각파")
    add_sentences(doc, [
        "칠판에 적힌 다음 입력은 삼각파다.",
        "주파수는 3 kHz다.",
        "진폭은 칠판의 전원 표시 10 Vpp를 썼다.",
        "15 kΩ과 100 nF만 쓴 이론 출력은 포물선이다.",
        "그 pk-pk는 0.2778 Vpp다.",
        "33 kΩ을 병렬로 넣은 삼각파 계산은 입력 없음.",
        "스코프 화면의 색, 채널, pk-pk, 시간축, 전압축, 촬영 시각은 입력 없음.",
    ])
    paragraph(doc, "[표 5-2] 삼각파", size=10, align="center", before=4, after=2)
    add_table(
        doc,
        ["항목", "값"],
        [
            ["칠판 입력", "삼각파, 10 Vpp, 3 kHz"],
            ["스코프 pk-pk", "입력 없음"],
            ["시간축", "입력 없음"],
            ["전압축", "입력 없음"],
            ["이론 출력 (15 kΩ, 100 nF)", "0.2778 Vpp (계산)"],
            ["이론 출력 (33 kΩ 병렬)", "입력 없음"],
        ],
        [62, 98],
    )

    subhead(doc, "5-3. 정리")
    paragraph(doc, "[표 5-3] 사인파와 삼각파", size=10, align="center", before=2, after=2)
    add_table(
        doc,
        ["구분", "입력", "출력"],
        [
            ["사인파", "10 Vpp, 3 kHz", "0.3537 Vpp (계산)"],
            ["사인파, 33 kΩ 병렬", "10 Vpp, 3 kHz", "0.3536 Vpp (계산)"],
            ["삼각파", "10 Vpp, 3 kHz", "0.2778 Vpp (계산)"],
            ["스코프", "입력 없음", "입력 없음"],
        ],
        [48, 46, 66],
    )

    heading(doc, "6. 결과 분석 및 고찰")
    subhead(doc, "6-1. 입력 없음")
    add_sentences(doc, [
        "스코프 화면이 입력에 없다.",
        "화면에서 이상해 보인 점은 입력 없음.",
    ])

    subhead(doc, "6-2. 크기")
    add_sentences(doc, [
        "비교할 측정값이 입력에 없다.",
        "오차율은 입력 없음.",
    ])

    subhead(doc, "6-3. 느낀점")
    add_sentences(doc, [
        "입력 없음.",
    ])

    heading(doc, "7. 결론")
    add_sentences(doc, [
        "15 kΩ, 100 nF, 33 kΩ 적분기의 칠판 입력은 3 kHz, 10 Vpp다.",
        "사인파 계산 출력은 0.3537 Vpp이고, 33 kΩ 병렬을 넣은 계산은 0.3536 Vpp다.",
        "삼각파 계산 출력은 0.2778 Vpp인 포물선이고, 스코프에서 읽은 출력과 오차율은 입력 없음.",
    ])

    heading(doc, "8. 참고문헌")
    add_sentences(doc, [
        "[1] 수업 화면, 실험 4 적분기.",
    ])

    out = ROOT / "20234092_이민규_회로이론실습설계2_실험4_적분기.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    build()
